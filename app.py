import datetime
import io
import os
import shutil
import traceback
import pandas as pd
import psycopg2
import streamlit as st
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- KİŞİSEL İMZA & KİMLİK ---
GELISTIRICI_ADI = "Ufuk Atak"

# Sayfa Yapılandırması
st.set_page_config(
    page_title="Sistem ve Parça Takip Sistemi", page_icon="🚀", layout="wide"
)

try:
    UPLOAD_FOLDER = "yuklenen_dosyalar"
    AUTOSAVE_FOLDER = "otomatik_yedekler"

    for folder in [UPLOAD_FOLDER, AUTOSAVE_FOLDER]:
        os.makedirs(folder, exist_ok=True)


    # --- POSTGRESQL (SUPABASE) BAĞLANTI FONKSİYONU ---
    def get_db_connection():
        return psycopg2.connect(st.secrets["postgres"]["connection_string"])


    # --- Önbellekli Veri Okuma Fonksiyonları ---
    @st.cache_data
    def veri_cek_parcalar():
        with get_db_connection() as conn:
            return pd.read_sql_query("SELECT * FROM parcalar", conn)


    @st.cache_data
    def veri_cek_notlar():
        with get_db_connection() as conn:
            return pd.read_sql_query(
                "SELECT * FROM gunluk_notlar ORDER BY id DESC", conn
            )


    @st.cache_data
    def veri_cek_arizalar():
        with get_db_connection() as conn:
            return pd.read_sql_query(
                "SELECT * FROM arizali_birimler ORDER BY id DESC", conn
            )


    @st.cache_data
    def veri_cek_loglar():
        with get_db_connection() as conn:
            return pd.read_sql_query(
                "SELECT * FROM islem_loglari ORDER BY id DESC", conn
            )


    @st.cache_data
    def veri_cek_depo():
        with get_db_connection() as conn:
            return pd.read_sql_query(
                "SELECT * FROM depo_stok ORDER BY parca_adi ASC", conn
            )


    # --- TÜRKÇE BÜYÜK HARF DÖNÜŞÜMÜ ---
    def tr_upper(text):
        if not isinstance(text, str):
            return str(text) if text is not None else ""
        return (
            text.replace("i", "İ")
            .replace("ı", "I")
            .replace("ş", "Ş")
            .replace("ğ", "Ğ")
            .replace("ü", "Ü")
            .replace("ö", "Ö")
            .replace("ç", "Ç")
            .upper()
        )


    # --- MODERN EXCEL TABLO BİÇİMLENDİRME FONKSİYONU ---
    def excel_stili_uygula(writer):
        """Verileri profesyonel Excel Tablosu (TableStyle) formatına dönüştürür ve sütunları otomatik ayarlar."""
        for sheetname in writer.sheets:
            ws = writer.sheets[sheetname]
            if ws.max_row > 1 and ws.max_column > 0:
                max_col_letter = get_column_letter(ws.max_column)
                table_range = f"A1:{max_col_letter}{ws.max_row}"

                clean_name = (
                    sheetname.replace(" ", "_")
                    .replace("İ", "I")
                    .replace("ğ", "g")
                    .replace("ü", "u")
                    .replace("ş", "s")
                    .replace("ö", "o")
                    .replace("ç", "c")
                )

                tab = Table(displayName=f"Tablo_{clean_name}", ref=table_range)

                style = TableStyleInfo(
                    name="TableStyleMedium9",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                tab.tableStyleInfo = style
                ws.add_table(tab)

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 5, 15)


    # --- OTOMATİK EXCEL YEDEKLEME FONKSİYONU ---
    def otomatik_yedekle():
        try:
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = os.path.join(
                AUTOSAVE_FOLDER, f"autosave_{timestamp}.xlsx"
            )
            latest_backup = os.path.join(
                AUTOSAVE_FOLDER, "sistem_takip_latest_autosave.xlsx"
            )

            df_p = veri_cek_parcalar()
            df_n = veri_cek_notlar()
            df_a = veri_cek_arizalar()
            df_l = veri_cek_loglar()
            df_d = veri_cek_depo()
            with get_db_connection() as conn:
                df_g = pd.read_sql_query("SELECT * FROM parca_gecmis", conn)

            with pd.ExcelWriter(backup_path, engine="openpyxl") as writer:
                df_p.to_excel(writer, sheet_name="Parcalar", index=False)
                df_n.to_excel(writer, sheet_name="Gunluk_Notlar", index=False)
                df_a.to_excel(writer, sheet_name="Arizali_Birimler", index=False)
                df_l.to_excel(writer, sheet_name="Loglar", index=False)
                df_d.to_excel(writer, sheet_name="Depo_Stok", index=False)
                df_g.to_excel(writer, sheet_name="Parca_Gecmisi", index=False)
                excel_stili_uygula(writer)

            shutil.copy2(backup_path, latest_backup)

            yedekler = sorted(
                [
                    os.path.join(AUTOSAVE_FOLDER, f)
                    for f in os.listdir(AUTOSAVE_FOLDER)
                    if f.startswith("autosave_") and f.endswith(".xlsx")
                ],
                key=os.path.getmtime,
            )
            if len(yedekler) > 10:
                for eski_yedek in yedekler[:-10]:
                    os.remove(eski_yedek)
        except Exception as ex:
            print(f"Otomatik Excel yedekleme hatası: {ex}")


    # --- TARİH DÖNÜŞÜM YARDIMCILARI ---
    def str_to_date(tarih_str):
        if not tarih_str or pd.isna(tarih_str):
            return datetime.date.today()
        try:
            return datetime.datetime.strptime(
                str(tarih_str).strip(), "%d.%m.%Y"
            ).date()
        except ValueError:
            return datetime.date.today()


    # --- VERİTABANI BAŞLATMA VE MIGRATION ---
    def veritabanini_hazirla():
        with get_db_connection() as conn:
            with conn.cursor() as cursor:
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS parcalar (id SERIAL PRIMARY KEY)"
                )
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS islem_loglari (id SERIAL PRIMARY KEY, zaman TEXT, islem_turu TEXT, detay TEXT)"
                )
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS gunluk_notlar (id SERIAL PRIMARY KEY, tarih TEXT, bolge TEXT, detay TEXT)"
                )
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS arizali_birimler (id SERIAL PRIMARY KEY, tarih TEXT, bolge TEXT, bildirim TEXT)"
                )
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS parca_gecmis (id SERIAL PRIMARY KEY, parca_sn TEXT, tarih TEXT, islem TEXT, aciklama TEXT)"
                )
                cursor.execute(
                    "CREATE TABLE IF NOT EXISTS depo_stok (id SERIAL PRIMARY KEY, bolge TEXT, sistem_adi TEXT, sistem_pn TEXT, sistem_sn TEXT, parca_adi TEXT, parca_pn TEXT, parca_sn TEXT, durum TEXT, ekleme_tarihi TEXT, aciklama TEXT)"
                )
            conn.commit()


    def log_yaz(islem_turu, detay):
        try:
            zaman = datetime.datetime.now().strftime("%d.%m.%Y %H:%M:%S")
            with get_db_connection() as conn:
                with conn.cursor() as cursor:
                    cursor.execute(
                        "INSERT INTO islem_loglari (zaman, islem_turu, detay) VALUES (%s, %s, %s)",
                        (zaman, islem_turu, detay),
                    )
                conn.commit()
            otomatik_yedekle()
        except Exception as e:
            print(f"Log yazılamadı: {e}")


    veritabanini_hazirla()

    # --- ÖZEL SİYAH TEMA, ÇERÇEVE VE ENVANTER KART CSS ---
    st.markdown(
        """
        <style>
        .stApp {
            background-color: #121212;
            color: #e0e0e0;
        }
        .metric-card {
            background-color: #1e1e1e;
            border: 1px solid #333333;
            border-radius: 8px;
            padding: 12px 15px;
            text-align: center;
            box-shadow: 0 2px 4px rgba(0,0,0,0.2);
        }
        .metric-title {
            font-size: 13px;
            color: #a0a0a0;
            font-weight: 600;
            margin-bottom: 4px;
            text-transform: uppercase;
        }
        .metric-value {
            font-size: 22px;
            color: #ffffff;
            font-weight: 700;
        }
        .autosave-badge {
            font-size: 12px;
            color: #2ecc71;
            background-color: #1b382b;
            padding: 4px 8px;
            border-radius: 4px;
            border: 1px solid #2ecc71;
        }
        .top-signature {
            background: #1e1e1e;
            border: 1px solid #333333;
            color: #a0a0a0;
            padding: 6px 14px;
            border-radius: 20px;
            font-size: 12px;
            letter-spacing: 0.5px;
            text-align: right;
            display: inline-block;
            float: right;
            box-shadow: 0 2px 4px rgba(0,0,0,0.2);
        }
        .top-signature b {
            color: #00adb5;
        }
        .item-card {
            background-color: #1a1a1a;
            border: 1px solid #2d2d2d;
            border-left: 4px solid #00adb5;
            border-radius: 6px;
            padding: 12px 16px;
            margin-bottom: 10px;
        }
        .item-title {
            font-size: 16px;
            font-weight: bold;
            color: #ffffff;
            margin-bottom: 4px;
        }
        .item-sub {
            font-size: 13px;
            color: #b0b0b0;
        }
        .tag-badge {
            display: inline-block;
            padding: 3px 8px;
            border-radius: 4px;
            font-size: 11px;
            font-weight: 600;
            margin-right: 5px;
        }
        .status-faal { background-color: #1e3d2f; color: #2ecc71; border: 1px solid #2ecc71; }
        .status-onarim { background-color: #3d351e; color: #f1c40f; border: 1px solid #f1c40f; }
        .status-gayri { background-color: #3d1e1e; color: #e74c3c; border: 1px solid #e74c3c; }
        .status-yedek { background-color: #1e2f3d; color: #3498db; border: 1px solid #3498db; }
        .status-sevk { background-color: #2e1e3d; color: #9b59b6; border: 1px solid #9b59b6; }
        
        .region-header {
            background: linear-gradient(90deg, #1e2a38 0%, #1a1a1a 100%);
            border: 1px solid #00adb5;
            border-radius: 8px;
            padding: 10px 15px;
            margin: 15px 0 10px 0;
            color: #00adb5;
            font-size: 18px;
            font-weight: bold;
        }
        .system-header {
            background-color: #222222;
            border-left: 3px solid #f1c40f;
            padding: 6px 12px;
            margin: 8px 0;
            border-radius: 4px;
            color: #ffffff;
            font-weight: 600;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    # --- ÜST BAŞLIK, SUNUCU DURUMU VE SAĞ ÜST İMZA ---
    col_head1, col_head2, col_head3 = st.columns([4, 3, 3])
    with col_head1:
        st.markdown("### 🚀 Sistem ve Parça Takip")
    with col_head2:
        st.markdown(
            "<div style='padding-top: 8px;'><span class='autosave-badge'>🟢 Stockholm Sunucu Aktif</span></div>",
            unsafe_allow_html=True,
        )
    with col_head3:
        st.markdown(
            f"<div style='text-align: right;'><div class='top-signature'>Design by <b>{GELISTIRICI_ADI}</b> ⚡</div></div>",
            unsafe_allow_html=True,
        )

    df_parcalar = veri_cek_parcalar()
    df_depo = veri_cek_depo()

    toplam = len(df_parcalar)
    faal = len(df_parcalar[df_parcalar["durum"] == "FAAL"]) if not df_parcalar.empty and "durum" in df_parcalar.columns else 0
    yedek = len(df_parcalar[df_parcalar["durum"] == "YEDEK PARÇA"]) if not df_parcalar.empty and "durum" in df_parcalar.columns else 0
    onarimda = len(df_parcalar[df_parcalar["durum"] == "ONARIMDA"]) if not df_parcalar.empty and "durum" in df_parcalar.columns else 0
    gayri = len(df_parcalar[df_parcalar["durum"] == "GAYRI FAAL"]) if not df_parcalar.empty and "durum" in df_parcalar.columns else 0

    kritik = 0
    kritik_liste = []
    if not df_parcalar.empty and "onarim_tarih" in df_parcalar.columns and "durum" in df_parcalar.columns:
        bugun = datetime.date.today()
        for _, row in df_parcalar[df_parcalar["durum"] == "ONARIMDA"].iterrows():
            o_tarih = row["onarim_tarih"]
            if o_tarih:
                try:
                    baslangic = datetime.datetime.strptime(str(o_tarih).strip(), "%d.%m.%Y").date()
                    gecen_gun = (bugun - baslangic).days
                    if gecen_gun >= 30:
                        kritik += 1
                        kritik_liste.append(
                            f"• **Bölge: {row.get('bolge', 'Bölge Yok')}** ({row.get('parça_adi', row.get('parca_adi', 'Parça'))} - SN: {row.get('parca_sn', '-')}) -> {gecen_gun} gündür onarımda!"
                        )
                except ValueError:
                    pass

    if kritik > 0:
        st.error(f"⚠️ **DİKKAT:** 30 Günü Aşan Onarımda Bekleyen **{kritik}** Adet Parça Bulunuyor!")
        with st.expander("Kritik Parçaları Listele", expanded=False):
            for k_bilgi in kritik_liste:
                st.markdown(k_bilgi)

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    with c1:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Toplam</div><div class="metric-value">{toplam}</div></div>', unsafe_allow_html=True)
    with c2:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Faal</div><div class="metric-value" style="color: #2ecc71;">{faal}</div></div>', unsafe_allow_html=True)
    with c3:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Yedek</div><div class="metric-value" style="color: #3498db;">{yedek}</div></div>', unsafe_allow_html=True)
    with c4:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Onarımda</div><div class="metric-value" style="color: #f1c40f;">{onarimda}</div></div>', unsafe_allow_html=True)
    with c5:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Kritik (30+)</div><div class="metric-value" style="color: #e74c3c;">{kritik}</div></div>', unsafe_allow_html=True)
    with c6:
        st.markdown(f'<div class="metric-card"><div class="metric-title">Gayri Faal</div><div class="metric-value" style="color: #95a5a6;">{gayri}</div></div>', unsafe_allow_html=True)

    st.markdown("<hr style='margin:15px 0 10px 0; border-color: #333333;'>", unsafe_allow_html=True)

    tab_takip, tab_ekle, tab_arizali, tab_notlar, tab_loglar, tab_gecmis, tab_yonetim, tab_depo = st.tabs([
        "📊 Envanter & Durum Takip",
        "➕ Yeni Kayıt",
        "⚠️ Arızalı Birimler",
        "📝 Günlük Notlar",
        "📜 Loglar",
        "🔍 Parça Geçmişi",
        "⚙️ Yönetim",
        "📦 Depo / Sevk Listesi",
    ])

    @st.dialog("Kayıt Düzenle", width="large")
    def duzenle_dialog(r_id):
        with get_db_connection() as conn:
            df_tekil = pd.read_sql_query("SELECT * FROM parcalar WHERE id = %s", conn, params=(r_id,))

        if not df_tekil.empty:
            row = df_tekil.iloc[0]
            with st.form(key=f"dialog_form_{r_id}"):
                e_bolge = st.text_input("Bölge", value=(row.get("bolge", "") if pd.notna(row.get("bolge", "")) else ""))
                e_sistem = st.text_input("Sistem Adı", value=(row.get("sistem_adi", "") if pd.notna(row.get("sistem_adi", "")) else ""))

                col_d1, col_d2 = st.columns(2)
                e_s_pn = col_d1.text_input("Sistem PN", value=(row.get("sistem_pn", "") if pd.notna(row.get("sistem_pn", "")) else ""))
                e_s_sn = col_d2.text_input("Sistem SN", value=(row.get("sistem_sn", "") if pd.notna(row.get("sistem_sn", "")) else ""))

                e_parca = st.text_input("Parça Adı", value=(row.get("parca_adi", "") if pd.notna(row.get("parca_adi", "")) else ""))

                col_d3, col_d4 = st.columns(2)
                e_p_pn = col_d3.text_input("Parça PN", value=(row.get("parca_pn", "") if pd.notna(row.get("parca_pn", "")) else ""))
                e_p_sn = col_d4.text_input("Parça SN", value=(row.get("parca_sn", "") if pd.notna(row.get("parca_sn", "")) else ""))

                mevcut_d = row.get("durum", "FAAL")
                d_list = ["FAAL", "YEDEK PARÇA", "ONARIMDA", "SEVK 🚚", "GAYRI FAAL"]
                d_idx = d_list.index(mevcut_d) if mevcut_d in d_list else 0
                e_durum = st.selectbox("Durum", d_list, index=d_idx)

                varsayilan_tarih = str_to_date(row.get("onarim_tarih"))
                e_tarih_obj = st.date_input("Onarım Tarihi", value=varsayilan_tarih, format="DD.MM.YYYY")
                e_tarih_str = e_tarih_obj.strftime("%d.%m.%Y")

                e_aciklama = st.text_area("Açıklama / Not", value=(row.get("aciklama", "") if pd.notna(row.get("aciklama", "")) else ""))

                if st.form_submit_button("Güncellemeyi Kaydet", type="primary"):
                    with get_db_connection() as conn:
                        with conn.cursor() as cursor:
                            cursor.execute(
                                """
                                UPDATE parcalar SET bolge=%s, sistem_adi=%s, sistem_pn=%s, sistem_sn=%s, parca_adi=%s, parca_pn=%s, parca_sn=%s, durum=%s, onarim_tarih=%s, aciklama=%s
                                WHERE id=%s
                                """,
                                (
                                    tr_upper(e_bolge),
                                    tr_upper(e_sistem),
                                    tr_upper(e_s_pn),
                                    tr_upper(e_s_sn),
                                    tr_upper(e_parca),
                                    tr_upper(e_p_pn),
                                    tr_upper(e_p_sn),
                                    e_durum,
                                    e_tarih_str,
                                    tr_upper(e_aciklama),
                                    r_id,
                                ),
                            )

                            if e_durum == "SEVK 🚚":
                                cursor.execute(
                                    """
                                    INSERT INTO depo_stok (bolge, sistem_adi, sistem_pn, sistem_sn, parca_adi, parca_pn, parca_sn, durum, ekleme_tarihi, aciklama)
                                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                    """,
                                    (
                                        tr_upper(e_bolge),
                                        tr_upper(e_sistem),
                                        tr_upper(e_s_pn),
                                        tr_upper(e_s_sn),
                                        tr_upper(e_parca),
                                        tr_upper(e_p_pn),
                                        tr_upper(e_p_sn),
                                        e_durum,
                                        e_tarih_str,
                                        tr_upper(e_aciklama),
                                    ),
                                )
                            cursor.execute(
                                "INSERT INTO parca_gecmis (parca_sn, tarih, islem, aciklama) VALUES (%s, %s, %s, %s)",
                                (
                                    tr_upper(e_p_sn),
                                    datetime.datetime.now().strftime("%d.%m.%Y %H:%M"),
                                    f"GÜNCELLEME ({e_durum})",
                                    tr_upper(e_aciklama),
                                ),
                            )
                        conn.commit()

                    st.cache_data.clear()
                    log_yaz("GÜNCELLEME", f"ID {r_id} güncellendi.")
                    st.toast("✅ Kayıt güncellendi!", icon="💾")
                    st.rerun()

    @st.dialog("Silme Onayı")
    def sil_onay_dialog(r_id, parca_adi, parca_sn):
        st.warning(f"**{parca_adi}** (SN: `{parca_sn}`) kaydını silmek istediğinizden emin misiniz?")
        col_s1, col_s2 = st.columns(2)
        if col_s1.button("Evet, Sil", type="primary", use_container_width=True):
            with get_db_connection() as conn:
                with conn.cursor() as cursor:
                    cursor.execute("DELETE FROM parcalar WHERE id=%s", (r_id,))
                conn.commit()
            st.cache_data.clear()
            log_yaz("SİLME", f"ID {r_id} ({parca_adi}) silindi.")
            st.toast("✅ Kayıt silindi.", icon="🗑️")
            st.rerun()
        if col_s2.button("İptal", use_container_width=True):
            st.rerun()

    # 1. SEKME: ENVANTER & BÖLGE-BİRİM ÇİFT GRUPLAMALI DURUM TAKİP
    with tab_takip:
        if not df_parcalar.empty:
            df_envanter = df_parcalar[df_parcalar["durum"] != "SEVK 🚚"].copy()

            col_f1, col_f2, col_f3, col_f4 = st.columns(4)

            bolgeler = ["TÜMÜ"] + sorted(list(set([tr_upper(b).strip() for b in df_envanter["bolge"].dropna() if str(b).strip()]))) if "bolge" in df_envanter.columns else ["TÜMÜ"]
            parcalar = ["TÜMÜ"] + sorted(list(set([tr_upper(p).strip() for p in df_envanter["parca_adi"].dropna() if str(p).strip()]))) if "parca_adi" in df_envanter.columns else ["TÜMÜ"]
            durumlar = ["TÜMÜ", "FAAL", "YEDEK PARÇA", "ONARIMDA", "GAYRI FAAL", "30 GÜN+ KRİTİK"]

            f_bolge = col_f1.selectbox("📍 Bölge Filtresi (Şırnak, Cizre, Mardin vb.)", bolgeler)
            f_parca = col_f2.selectbox("📦 Parça Filtresi", parcalar)
            f_durum = col_f3.selectbox("📊 Durum Filtresi", durumlar)
            f_arama = col_f4.text_input("🔍 Hızlı Arama")

            filt_df = df_envanter.copy()
            
            if f_bolge != "TÜMÜ" and "bolge" in filt_df.columns:
                filt_df = filt_df[filt_df["bolge"].apply(lambda x: tr_upper(x).strip() == f_bolge)]
            if f_parca != "TÜMÜ" and "parca_adi" in filt_df.columns:
                filt_df = filt_df[filt_df["parca_adi"].apply(lambda x: tr_upper(x).strip() == f_parca)]
            if f_durum != "TÜMÜ" and f_durum != "30 GÜN+ KRİTİK" and "durum" in filt_df.columns:
                filt_df = filt_df[filt_df["durum"] == f_durum]
            elif f_durum == "30 GÜN+ KRİTİK" and "durum" in filt_df.columns and "onarim_tarih" in filt_df.columns:
                kritik_idler = []
                bugun = datetime.date.today()
                for _, r in filt_df[filt_df["durum"] == "ONARIMDA"].iterrows():
                    if r["onarim_tarih"]:
                        try:
                            b_t = datetime.datetime.strptime(str(r["onarim_tarih"]).strip(), "%d.%m.%Y").date()
                            if (bugun - b_t).days >= 30:
                                kritik_idler.append(r["id"])
                        except ValueError:
                            pass
                filt_df = filt_df[filt_df["id"].isin(kritik_idler)]

            if f_arama:
                a_upper = tr_upper(f_arama)
                filt_df = filt_df[
                    filt_df.apply(lambda row: row.astype(str).str.upper().str.contains(a_upper).any(), axis=1)
                ]

            st.markdown("<hr style='margin:10px 0 15px 0; border-color: #333;'>", unsafe_allow_html=True)

            if not filt_df.empty and "bolge" in filt_df.columns:
                b_gruplari = filt_df.groupby("bolge")
                
                st.markdown(f"##### 📍 Toplam **{len(b_gruplari)}** Bölge Altındaki Sistemler ve Birimler Gruplandırıldı")
                
                # BÖLGE DÖNGÜSÜ
                for b_adi, b_grubu in b_gruplari:
                    b_toplam = len(b_grubu)
                    b_faal = len(b_grubu[b_grubu["durum"] == "FAAL"])
                    b_onarim = len(b_grubu[b_grubu["durum"] == "ONARIMDA"])
                    
                    durum_ozet = f"🟢 {b_faal} Faal"
                    if b_onarim > 0:
                        durum_ozet += f" | 🟡 {b_onarim} Onarımda"

                    # BÖLGE BAŞLIĞI
                    st.markdown(
                        f"""
                        <div class="region-header">
                            📍 {b_adi} BÖLGESİ — <span style='font-size: 14px; color: #a0a0a0;'>Toplam: {b_toplam} Malzeme/Parça ({durum_ozet})</span>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )

                    # BÖLGE İÇİNDEKİ SİSTEMLERİ/BİRİMLERİ GRUPLA
                    sistem_gruplari = b_grubu.groupby("sistem_adi")
                    
                    for s_adi, s_grubu in sistem_gruplari:
                        s_sn_ornek = s_grubu["sistem_sn"].iloc[0] if "sistem_sn" in s_grubu.columns and pd.notna(s_grubu["sistem_sn"].iloc[0]) else "-"
                        
                        with st.expander(f"🖥️ {s_adi} (Sistem SN: {s_sn_ornek}) — [{len(s_grubu)} Adet Parça]", expanded=True):
                            for _, row in s_grubu.iterrows():
                                r_id = row["id"]
                                durum = row.get("durum", "FAAL")
                                p_sn = row.get("parca_sn", "-")
                                p_pn = row.get("parca_pn", "-")
                                p_adi = row.get("parca_adi", "Parça")

                                status_class = (
                                    "status-faal" if durum == "FAAL"
                                    else ("status-onarim" if durum == "ONARIMDA"
                                    else ("status-gayri" if durum == "GAYRI FAAL"
                                    else ("status-yedek" if durum == "YEDEK PARÇA" else "status-sevk")))
                                )

                                c_card, c_acts = st.columns([8.5, 1.5])
                                with c_card:
                                    aciklama_html = f"<div style='font-size:12px; color:#888; margin-top:4px;'>📝 Not: {row.get('aciklama')}</div>" if pd.notna(row.get("aciklama")) and str(row.get("aciklama")).strip() else ""
                                    st.markdown(
                                        f"""
                                        <div class="item-card">
                                            <div class="item-title">📦 {p_adi} <span class="tag-badge {status_class}">{durum}</span></div>
                                            <div class="item-sub">
                                                <b>Parça SN:</b> <code style='color:#00adb5;'>{p_sn}</code> | <b>PN:</b> {p_pn} | <b>Tarih:</b> {row.get('onarim_tarih', '-')}
                                            </div>
                                            {aciklama_html}
                                        </div>
                                        """,
                                        unsafe_allow_html=True,
                                    )
                                with c_acts:
                                    col_b1, col_b2 = st.columns(2)
                                    with col_b1:
                                        if st.button("✏️", key=f"edit_{r_id}", help="Düzenle", use_container_width=True):
                                            duzenle_dialog(r_id)
                                    with col_b2:
                                        if st.button("🗑️", key=f"del_{r_id}", help="Sil", use_container_width=True):
                                            sil_onay_dialog(r_id, p_adi, p_sn)
            else:
                st.info("Filtreleme kriterlerine uygun kayıt bulunamadı.")
        else:
            st.info("Kayıt bulunmuyor.")

    # 2. SEKME: YENİ KAYIT EKLE
    with tab_ekle:
        kayitli_bolgeler = sorted(list(set([tr_upper(b).strip() for b in df_parcalar["bolge"].dropna() if str(b).strip()]))) if not df_parcalar.empty and "bolge" in df_parcalar.columns else []

        st.markdown("#### ➕ Yeni Sistem ve Parça Kaydı")

        st.markdown("##### 📍 Bölge Seçim Tipi")
        secim_tipi = st.radio(
            "Bölge Giriş Tipi", 
            ["Kayıtlı Bölge", "Yeni Bölge"], 
            horizontal=True, 
            label_visibility="collapsed"
        )

        with st.form("yeni_kayit_formu", clear_on_submit=False):
            form_col_left, form_col_right = st.columns(2)

            # SOL SÜTUN: BÖLGE VE CİHAZ BİLGİLERİ
            with form_col_left:
                st.markdown("#### 📍 Bölge & Cihaz Bilgileri")
                
                aktif_bolge = ""
                if secim_tipi == "Kayıtlı Bölge" and len(kayitli_bolgeler) > 0:
                    aktif_bolge = st.selectbox("Bölge Seç", kayitli_bolgeler)
                else:
                    aktif_bolge = st.text_input("Yeni Bölge Adı (Örn: Şırnak Merkez, Cizre, Mardin)")

                varsayilan_sistem = ""
                varsayilan_s_pn = ""
                varsayilan_s_sn = ""

                if aktif_bolge and not df_parcalar.empty and "bolge" in df_parcalar.columns:
                    eslesen_kayitlar = df_parcalar[df_parcalar["bolge"].str.upper() == str(aktif_bolge).upper()]
                    if not eslesen_kayitlar.empty:
                        son_kayit = eslesen_kayitlar.iloc[-1]
                        varsayilan_sistem = son_kayit.get("sistem_adi", "") if pd.notna(son_kayit.get("sistem_adi", "")) else ""
                        varsayilan_s_pn = son_kayit.get("sistem_pn", "") if pd.notna(son_kayit.get("sistem_pn", "")) else ""
                        varsayilan_s_sn = son_kayit.get("sistem_sn", "") if pd.notna(son_kayit.get("sistem_sn", "")) else ""

                e_sistem = st.text_input("Sistem Adı", value=varsayilan_sistem)
                sub_c1, sub_c2 = st.columns(2)
                e_s_pn = sub_c1.text_input("Sistem PN", value=varsayilan_s_pn)
                e_s_sn = sub_c2.text_input("Sistem SN", value=varsayilan_s_sn)

            # SAĞ SÜTUN: PARÇA BİLGİLERİ
            with form_col_right:
                st.markdown("#### 📦 Parça Detayları")
                
                e_parca = st.text_input("Parça Adı (Yazıp Enter'a basın)", key="form_parca_adi_giris")
                
                bulunan_pn = ""
                if e_parca and not df_parcalar.empty and "parca_adi" in df_parcalar.columns:
                    eslesen_parcalar = df_parcalar[df_parcalar["parca_adi"].str.upper() == str(e_parca).strip().upper()]
                    if not eslesen_parcalar.empty:
                        p_son = eslesen_parcalar.iloc[-1]
                        p_val = p_son.get("parca_pn", "")
                        if pd.notna(p_val):
                            bulunan_pn = str(p_val)

                sub_c3, sub_c4 = st.columns(2)
                e_p_pn = sub_c3.text_input("Parça PN (Otomatik / Sabit)", value=bulunan_pn)
                e_p_sn = sub_c4.text_input("Parça SN (Her Parça İçin Boş / Değişken)", value="", key="form_parca_sn_giris")

                e_durum = st.selectbox("Durum", ["FAAL", "YEDEK PARÇA", "ONARIMDA", "SEVK 🚚", "GAYRI FAAL"])
                
                e_tarih_obj = st.date_input("Onarım / Ekleme Tarihi", value=datetime.date.today(), format="DD.MM.YYYY")
                e_tarih_str = e_tarih_obj.strftime("%d.%m.%Y")

            st.markdown("<br>", unsafe_allow_html=True)
            
            e_aciklama = st.text_area("Açıklama / Not")
            yuklenen_dosya_form = st.file_uploader("Belge/Fotoğraf Yükle", type=["png", "jpg", "jpeg", "pdf"])

            st.markdown("<br>", unsafe_allow_html=True)
            if st.form_submit_button("🚀 Sisteme Kaydet", type="primary", use_container_width=True):
                if not aktif_bolge or not e_sistem or not e_parca:
                    st.warning("Bölge, Sistem Adı ve Parça Adı zorunludur!")
                else:
                    kayit_hash = f"{tr_upper(aktif_bolge)}_{tr_upper(e_sistem)}_{tr_upper(e_parca)}_{tr_upper(e_p_sn)}"
                    if st.session_state.get("son_islem_hash") == kayit_hash:
                        st.warning("⚠️ Bu kayıt az önce zaten eklendi!")
                    else:
                        st.session_state["son_islem_hash"] = kayit_hash

                        dosya_ismi = None
                        if yuklenen_dosya_form is not None:
                            dosya_ismi = f"{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}_{yuklenen_dosya_form.name}"
                            with open(os.path.join(UPLOAD_FOLDER, dosya_ismi), "wb") as f:
                                f.write(yuklenen_dosya_form.getbuffer())

                        with get_db_connection() as conn:
                            with conn.cursor() as cursor:
                                cursor.execute(
                                    """
                                    INSERT INTO parcalar (bolge, sistem_adi, sistem_pn, sistem_sn, parca_adi, parca_pn, parca_sn, durum, onarim_tarih, aciklama, dosya_adi)
                                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                    """,
                                    (
                                        tr_upper(aktif_bolge),
                                        tr_upper(e_sistem),
                                        tr_upper(e_s_pn),
                                        tr_upper(e_s_sn),
                                        tr_upper(e_parca),
                                        tr_upper(e_p_pn),
                                        tr_upper(e_p_sn),
                                        e_durum,
                                        e_tarih_str,
                                        tr_upper(e_aciklama),
                                        dosya_ismi,
                                    ),
                                )

                                if e_durum == "SEVK 🚚":
                                    cursor.execute(
                                        """
                                        INSERT INTO depo_stok (bolge, sistem_adi, sistem_pn, sistem_sn, parca_adi, parca_pn, parca_sn, durum, ekleme_tarihi, aciklama)
                                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                        """,
                                        (
                                            tr_upper(aktif_bolge),
                                            tr_upper(e_sistem),
                                            tr_upper(e_s_pn),
                                            tr_upper(e_s_sn),
                                            tr_upper(e_parca),
                                            tr_upper(e_p_pn),
                                            tr_upper(e_p_sn),
                                            e_durum,
                                            e_tarih_str,
                                            tr_upper(e_aciklama),
                                        ),
                                    )

                                cursor.execute(
                                    "INSERT INTO parca_gecmis (parca_sn, tarih, islem, aciklama) VALUES (%s, %s, %s, %s)",
                                    (
                                        tr_upper(e_p_sn),
                                        datetime.datetime.now().strftime("%d.%m.%Y %H:%M"),
                                        f"İLK KAYIT ({e_durum})",
                                        tr_upper(e_aciklama),
                                    ),
                                )
                            conn.commit()

                        st.cache_data.clear()
                        log_yaz("YENİ KAYIT", f"Bölge: {aktif_bolge}, Sistem: {e_sistem} eklendi.")
                        st.toast("✅ Kayıt başarıyla eklendi!", icon="🎉")
                        st.rerun()

    # 3. SEKME: ARIZALI BİRİMLER
    with tab_arizali:
        st.subheader("⚠️ Bölgesel Arıza Bildirimleri")
        st.markdown("Seri numarası veya parça kodu girmeden, doğrudan bölge seçip arıza bildirimi ekleyebilir veya çözülenleri listeden kaldırabilirsiniz.")

        kayitli_ariza_bolgeler = sorted(list(set([tr_upper(b).strip() for b in df_parcalar["bolge"].dropna() if str(b).strip()]))) if not df_parcalar.empty and "bolge" in df_parcalar.columns else []

        with st.form("hizli_ariza_form", clear_on_submit=True):
            st.markdown("##### ➕ Yeni Arıza Bildirimi Gir")
            a_secim_tipi = st.radio("Bölge Seçim Tipi", ["Kayıtlı Bölge Seç", "Yeni Bölge Yaz"], horizontal=True)
            if a_secim_tipi == "Kayıtlı Bölge Seç" and len(kayitli_ariza_bolgeler) > 0:
                a_bolge = st.selectbox("Bölge", kayitli_ariza_bolgeler)
            else:
                a_bolge = st.text_input("Bölge Adı")

            a_bildirim = st.text_area("Arıza Bildirimi (Örn: HSY arızası mevcut)")

            if st.form_submit_button("Bildirimi Kaydet", type="primary"):
                if not a_bolge or not a_bildirim:
                    st.warning("Bölge ve arıza bildirimi boş olamaz!")
                else:
                    tarih_str = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
                    with get_db_connection() as conn:
                        with conn.cursor() as cursor:
                            cursor.execute(
                                "INSERT INTO arizali_birimler (tarih, bolge, bildirim) VALUES (%s, %s, %s)",
                                (tarih_str, tr_upper(a_bolge), tr_upper(a_bildirim)),
                            )
                        conn.commit()

                    st.cache_data.clear()
                    log_yaz("ARIZA BİLDİRİMİ", f"Bölge: {a_bolge} - {a_bildirim}")
                    st.toast("✅ Arıza bildirimi eklendi!", icon="⚠️")
                    st.rerun()

        st.markdown("---")
        st.markdown("##### 📋 Aktif Arıza Bildirimleri Listesi")
        df_arizalar = veri_cek_arizalar()
        if not df_arizalar.empty:
            st.warning(f"Sistemde toplam **{len(df_arizalar)}** adet aktif arıza bildirimi bulunuyor.")
            for _, a_row in df_arizalar.iterrows():
                a_id = a_row["id"]
                col_b1, col_b2 = st.columns([8, 2])
                with col_b1:
                    st.markdown(
                        f"**📍 {a_row.get('bolge', '-')}** — <small style='color:#a0a0a0;'>{a_row.get('tarih', '-')}</small><br>"
                        f"🔴 {a_row.get('bildirim', '-')}",
                        unsafe_allow_html=True
                    )
                with col_b2:
                    if st.button("Çözüldü / Kaldır", key=f"ariza_sil_{a_id}", type="primary"):
                        with get_db_connection() as conn:
                            with conn.cursor() as cursor:
                                cursor.execute("DELETE FROM arizali_birimler WHERE id=%s", (a_id,))
                            conn.commit()
                        st.cache_data.clear()
                        log_yaz("ARIZA ÇÖZÜLDÜ", f"ID {a_id} numaralı arıza bildirimi kaldırıldı.")
                        st.toast("✅ Arıza kaldırıldı.", icon="🗑️")
                        st.rerun()
                st.markdown("<hr style='margin:5px 0; border-color: #333;'>", unsafe_allow_html=True)
        else:
            st.success("Harika! Aktif arıza bildirimi bulunmuyor. 🎉")

    # 4. SEKME: GÜNLÜK NOTLAR
    with tab_notlar:
        kayitli_not_bolgeler = sorted(list(set([tr_upper(b).strip() for b in df_parcalar["bolge"].dropna() if str(b).strip()]))) if not df_parcalar.empty and "bolge" in df_parcalar.columns else []
        st.subheader("📝 Günlük Notlar ve Olaylar")
        with st.form("gunluk_not_form", clear_on_submit=True):
            n_bolge = st.selectbox("İlgili Bölge / Konum", ["GENEL"] + kayitli_not_bolgeler)
            n_detay = st.text_area("Günlük Not / Faaliyet Açıklaması")
            if st.form_submit_button("Notu Kaydet", type="primary"):
                if not n_detay:
                    st.warning("Not içeriği boş olamaz!")
                else:
                    tarih_str = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
                    with get_db_connection() as conn:
                        with conn.cursor() as cursor:
                            cursor.execute(
                                "INSERT INTO gunluk_notlar (tarih, bolge, detay) VALUES (%s, %s, %s)",
                                (tarih_str, n_bolge, tr_upper(n_detay)),
                            )
                        conn.commit()
                    st.cache_data.clear()
                    log_yaz("GÜNLÜK NOT", f"Bölge: {n_bolge} için not eklendi.")
                    st.toast("✅ Not başarıyla kaydedildi!", icon="📝")
                    st.rerun()

        df_notlar = veri_cek_notlar()
        if not df_notlar.empty:
            st.markdown("---")
            for _, nrow in df_notlar.iterrows():
                st.markdown(f"**📅 {nrow.get('tarih', '-')} | 📍 {nrow.get('bolge', '-')}**")
                st.info(nrow.get("detay", ""))
        else:
            st.info("Henüz kaydedilmiş günlük not bulunmuyor.")

    # 5. SEKME: LOGLAR
    with tab_loglar:
        st.subheader("📜 Sistem İşlem Logları")
        df_loglar = veri_cek_loglar()
        if not df_loglar.empty:
            st.dataframe(df_loglar, use_container_width=True, hide_index=True)
            if st.button("Tüm Logları Temizle", type="primary"):
                with get_db_connection() as conn:
                    with conn.cursor() as cursor:
                        cursor.execute("DELETE FROM islem_loglari")
                    conn.commit()
                st.cache_data.clear()
                st.toast("🧹 Loglar temizlendi.")
                st.rerun()
        else:
            st.info("Kayıtlı log bulunmuyor.")

    # 6. SEKME: PARÇA GEÇMİŞİ
    with tab_gecmis:
        st.subheader("🔍 Parça Seri Numarası Geçmişi")
        s_ara_sn = st.text_input("Aranacak Parça Seri Numarası (SN)", key="parca_gecmis_input")
        if s_ara_sn:
            with get_db_connection() as conn:
                df_gcek = pd.read_sql_query(
                    "SELECT * FROM parca_gecmis WHERE parca_sn ILIKE %s ORDER BY id DESC",
                    conn,
                    params=(f"%{s_ara_sn.strip()}%",),
                )
            if not df_gcek.empty:
                st.success(f"`{s_ara_sn}` seri numaralı parça için {len(df_gcek)} adet işlem geçmişi bulundu:")
                for _, grow in df_gcek.iterrows():
                    st.markdown(f"• **{grow.get('tarih', '-')}** — `{grow.get('islem', '-')}`")
                    if pd.notna(grow.get("aciklama")) and str(grow.get("aciklama")).strip():
                        st.caption(f"Açıklama: {grow.get('aciklama')}")
                    st.markdown("---")
            else:
                st.info("Bu seri numarasına ait geçmiş işlem kaydı bulunamadı.")

    # 7. SEKME: YÖNETİM
    with tab_yonetim:
        st.subheader("⚙️ Veritabanı ve Yedek Yönetimi")
        col_y1, col_y2 = st.columns(2)

        with col_y1:
            st.markdown("##### 📥 Excel Yedek İndir")
            if st.button("Anlık Excel Raporu Oluştur ve İndir", type="primary"):
                output = io.BytesIO()
                df_p = veri_cek_parcalar()
                df_n = veri_cek_notlar()
                df_a = veri_cek_arizalar()
                df_l = veri_cek_loglar()
                df_d = veri_cek_depo()
                with get_db_connection() as conn:
                    df_g = pd.read_sql_query("SELECT * FROM parca_gecmis", conn)

                with pd.ExcelWriter(output, engine="openpyxl") as writer:
                    df_p.to_excel(writer, sheet_name="Parcalar", index=False)
                    df_n.to_excel(writer, sheet_name="Gunluk_Notlar", index=False)
                    df_a.to_excel(writer, sheet_name="Arizali_Birimler", index=False)
                    df_l.to_excel(writer, sheet_name="Loglar", index=False)
                    df_d.to_excel(writer, sheet_name="Depo_Stok", index=False)
                    df_g.to_excel(writer, sheet_name="Parca_Gecmisi", index=False)
                    excel_stili_uygula(writer)

                processed_data = output.getvalue()
                st.download_button(
                    label="💾 Excel Dosyasını İndir",
                    data=processed_data,
                    file_name=f"sistem_parca_takip_yedek_{datetime.date.today().strftime('%d_%m_%Y')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

        with col_y2:
            st.markdown("##### 📂 Otomatik Yedekler")
            yedek_dosyalari = sorted([f for f in os.listdir(AUTOSAVE_FOLDER) if f.endswith(".xlsx")], reverse=True)
            if yedek_dosyalari:
                secilen_yedek = st.selectbox("İndirilecek Otomatik Yedeği Seç", yedek_dosyalari)
                yedek_tam_yol = os.path.join(AUTOSAVE_FOLDER, secilen_yedek)
                with open(yedek_tam_yol, "rb") as yf:
                    st.download_button(
                        "⬇️ Seçilen Yedeği İndir",
                        data=yf,
                        file_name=secilen_yedek,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
            else:
                st.info("Henüz otomatik yedek bulunmuyor.")

    # 8. SEKME: DEPO / SEVK LİSTESİ
    with tab_depo:
        st.subheader("📦 Depo ve Sevk Yönetimi")
        st.markdown("Sevk edilen ve depodaki malzemeleri yönetebilir, listeleyebilirsiniz.")

        d_toplam = len(df_depo)
        d_konumlar = df_depo["bolge"].nunique() if not df_depo.empty and "bolge" in df_depo.columns else 0
        d_cesit = df_depo["parca_adi"].nunique() if not df_depo.empty and "parca_adi" in df_depo.columns else 0

        dc_m1, dc_m2, dc_m3, dc_m4 = st.columns([3, 3, 3, 3])
        dc_m1.metric("📦 Toplam Malzeme/Stok", d_toplam)
        dc_m2.metric("📍 Depo / Konum Sayısı", d_konumlar)
        dc_m3.metric("🏷️ Farklı Parça Çeşidi", d_cesit)
        
        with dc_m4:
            st.markdown("<div style='padding-top: 10px;'></div>", unsafe_allow_html=True)
            if st.button("🧹 İsimsiz Kayıtları Temizle", type="primary", use_container_width=True, help="Adı İSİMSİZ veya boş olan tüm depo kayıtlarını siler"):
                with get_db_connection() as conn:
                    with conn.cursor() as cursor:
                        cursor.execute("DELETE FROM depo_stok WHERE parca_adi ILIKE '%İSİMSİZ%' OR parca_adi = '' OR parca_adi IS NULL")
                    conn.commit()
                st.cache_data.clear()
                log_yaz("DEPO TEMİZLEME", "Hatalı ve İsimsiz depo kayıtları temizlendi.")
                st.toast("🧹 İsimsiz malzemeler depodan temizlendi!", icon="✨")
                st.rerun()

        st.markdown("<hr style='margin:15px 0; border-color: #333;'>", unsafe_allow_html=True)

        depo_islem_col1, depo_islem_col2 = st.columns(2)

        with depo_islem_col1:
            with st.expander("📥 Excel Raporu Çıkar", expanded=False):
                if not df_depo.empty:
                    output_depo = io.BytesIO()
                    with pd.ExcelWriter(output_depo, engine="openpyxl") as writer:
                        df_depo.to_excel(writer, sheet_name="Depo_Stok", index=False)
                        excel_stili_uygula(writer)
                    depo_excel_data = output_depo.getvalue()
                    
                    st.download_button(
                        label="💾 Depo Stok Raporunu İndir (.xlsx)",
                        data=depo_excel_data,
                        file_name=f"depo_stok_raporu_{datetime.date.today().strftime('%d_%m_%Y')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                else:
                    st.info("Depoda indirilecek malzeme bulunmuyor.")

        with depo_islem_col2:
            with st.expander("📤 Excel / CSV İle Toplu Malzeme Yükle", expanded=False):
                yuklenen_depo_excel = st.file_uploader("Depo Excel Dosyasını Seçin", type=["xlsx", "xls", "csv"], key="depo_excel_yukle")
                
                if yuklenen_depo_excel is not None:
                    try:
                        if yuklenen_depo_excel.name.endswith(".csv"):
                            df_yuklenen = pd.read_csv(yuklenen_depo_excel)
                        else:
                            df_yuklenen = pd.read_excel(yuklenen_depo_excel)
                        
                        st.success("✅ Excel dosyası başarıyla yüklendi! Lütfen sütunları eşleştirin:")
                        st.dataframe(df_yuklenen.head(3), use_container_width=True)
                        
                        kolonlar = list(df_yuklenen.columns)
                        
                        def tahmin_index(liste, anahtar_kelimeler):
                            for idx, col in enumerate(liste):
                                c_low = str(col).lower()
                                if any(k in c_low for k in anahtar_kelimeler):
                                    return idx
                            return 0

                        c_parca_idx = tahmin_index(kolonlar, ["tanim", "tanımı", "parca", "malzeme", "urun", "ad"])
                        c_pn_idx = tahmin_index(kolonlar, ["pn", "parca_pn", "kod"])
                        c_sn_idx = tahmin_index(kolonlar, ["sn", "parca_sn", "seri"])
                        c_not_idx = tahmin_index(kolonlar, ["not", "açıklama", "aciklama", "detay"])
                        c_bolge_idx = tahmin_index(kolonlar, ["bölge", "bolge", "konum", "depo"])

                        st.markdown("##### ⚙️ Sütun Eşleştirme Paneli")
                        map_col1, map_col2 = st.columns(2)
                        
                        with map_col1:
                            secilen_parca_col = st.selectbox("Parça Adı / Tanımı Sütunu *", kolonlar, index=c_parca_idx)
                            secilen_pn_col = st.selectbox("PN Sütunu (Opsiyonel)", ["Yok"] + kolonlar, index=(c_pn_idx + 1 if c_pn_idx < len(kolonlar) else 0))
                            secilen_sn_col = st.selectbox("SN Sütunu (Opsiyonel)", ["Yok"] + kolonlar, index=(c_sn_idx + 1 if c_sn_idx < len(kolonlar) else 0))
                        
                        with map_col2:
                            secilen_bolge_col = st.selectbox("Bölge / Konum Sütunu (Opsiyonel)", ["Sabit Merkez Depo"] + kolonlar, index=(c_bolge_idx + 1 if c_bolge_idx < len(kolonlar) else 0))
                            secilen_not_col = st.selectbox("Not / Açıklama Sütunu (Opsiyonel)", ["Yok"] + kolonlar, index=(c_not_idx + 1 if c_not_idx < len(kolonlar) else 0))

                        st.markdown("<br>", unsafe_allow_html=True)
                        if st.button("🚀 Verileri Depoya Aktarmaya Başla", type="primary", use_container_width=True):
                            eklenen_sayac = 0
                            tarih_str = datetime.date.today().strftime("%d.%m.%Y")
                            
                            with get_db_connection() as conn:
                                with conn.cursor() as cursor:
                                    for _, y_row in df_yuklenen.iterrows():
                                        b_parca = str(y_row[secilen_parca_col]).strip() if secilen_parca_col and pd.notna(y_row[secilen_parca_col]) else ""
                                        b_bolge = str(y_row[secilen_bolge_col]).strip() if secilen_bolge_col != "Sabit Merkez Depo" and pd.notna(y_row[secilen_bolge_col]) else "MERKEZ DEPO"
                                        b_pn = str(y_row[secilen_pn_col]).strip() if secilen_pn_col != "Yok" and pd.notna(y_row[secilen_pn_col]) else "-"
                                        b_sn = str(y_row[secilen_sn_col]).strip() if secilen_sn_col != "Yok" and pd.notna(y_row[secilen_sn_col]) else "-"
                                        b_not = str(y_row[secilen_not_col]).strip() if secilen_not_col != "Yok" and pd.notna(y_row[secilen_not_col]) else ""
                                        
                                        if b_parca and b_parca.lower() != "nan" and "isimsiz" not in b_parca.lower():
                                            cursor.execute(
                                                """
                                                INSERT INTO depo_stok (bolge, sistem_adi, sistem_pn, sistem_sn, parca_adi, parca_pn, parca_sn, durum, ekleme_tarihi, aciklama)
                                                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                                """,
                                                (
                                                    tr_upper(b_bolge),
                                                    "GENEL DEPO",
                                                    "-",
                                                    "-",
                                                    tr_upper(b_parca),
                                                    tr_upper(b_pn),
                                                    tr_upper(b_sn),
                                                    "SEVK 🚚",
                                                    tarih_str,
                                                    tr_upper(b_not),
                                                ),
                                            )
                                            eklenen_sayac += 1
                                    conn.commit()
                            
                            st.cache_data.clear()
                            log_yaz("TOPLU DEPO YÜKLEME", f"Excel'den {eklenen_sayac} adet malzeme depoya eklendi.")
                            st.success(f"🎉 Başarıyla {eklenen_sayac} malzeme depoya eklendi!")
                            st.rerun()
                    except Exception as ex:
                        st.error(f"Excel okunurken hata oluştu: {ex}")

        with st.expander("➕ Tekil Depo Malzemesi Ekle", expanded=False):
            with st.form("depoya_ekle_form", clear_on_submit=True):
                dc1, dc2 = st.columns(2)
                d_bolge_giris = dc1.text_input("Depo Konumu / Bölge (Örn: Merkez Depo)")
                d_sistem_giris = dc2.text_input("Sistem / Cihaz Adı (Opsiyonel)", value="GENEL DEPO")
                
                dc3, dc4 = st.columns(2)
                d_parca_giris = dc3.text_input("Malzeme / Parça Adı")
                d_pn_giris = dc4.text_input("Parça PN (Opsiyonel)", value="-")
                
                dc5, dc6 = st.columns(2)
                d_sn_giris = dc5.text_input("Parça SN (Opsiyonel)", value="-")
                d_not_giris = dc6.text_input("Açıklama / Adet Notu", value="")
                
                if st.form_submit_button("Depoya Malzeme Kaydet", type="primary"):
                    if not d_bolge_giris or not d_parca_giris:
                        st.warning("Bölge/Konum ve Malzeme Adı zorunludur!")
                    else:
                        tarih_str = datetime.date.today().strftime("%d.%m.%Y")
                        with get_db_connection() as conn:
                            with conn.cursor() as cursor:
                                cursor.execute(
                                    """
                                    INSERT INTO depo_stok (bolge, sistem_adi, sistem_pn, sistem_sn, parca_adi, parca_pn, parca_sn, durum, ekleme_tarihi, aciklama)
                                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                                    """,
                                    (
                                        tr_upper(d_bolge_giris),
                                        tr_upper(d_sistem_giris),
                                        "-",
                                        "-",
                                        tr_upper(d_parca_giris),
                                        tr_upper(d_pn_giris),
                                        tr_upper(d_sn_giris),
                                        "SEVK 🚚",
                                        tarih_str,
                                        tr_upper(d_not_giris),
                                    ),
                                )
                            conn.commit()
                        st.cache_data.clear()
                        log_yaz("DEPOYA EKLEME", f"Depoya dışarıdan {d_parca_giris} eklendi.")
                        st.toast("✅ Malzeme depoya başarıyla eklendi!", icon="📦")
                        st.rerun()

        st.markdown("<hr style='margin:15px 0; border-color: #333;'>", unsafe_allow_html=True)

        depo_arama = st.text_input("🔍 Depoda Malzeme Ara (Parça Adı, SN, PN veya Bölge)")

        if not df_depo.empty:
            filt_depo = df_depo.copy()
            if depo_arama:
                da_upper = tr_upper(depo_arama)
                filt_depo = filt_depo[
                    filt_depo.apply(lambda row: row.astype(str).str.upper().str.contains(da_upper).any(), axis=1)
                ]

            st.markdown(f"<small style='color:#a0a0a0;'>Listelenen Toplam Kayıt: {len(filt_depo)} | Farklı Malzeme Çeşidi: {filt_depo['parca_adi'].nunique() if not filt_depo.empty else 0}</small>", unsafe_allow_html=True)
            st.markdown("<br>", unsafe_allow_html=True)

            secilen_idler = []

            for parca_adi, grup in filt_depo.groupby("parca_adi", sort=False):
                bolgeler_str = ", ".join(grup["bolge"].dropna().unique())
                
                with st.expander(f"📦 {parca_adi}  ({len(grup)} Adet) — [Bölgeler: {bolgeler_str}]", expanded=False):
                    for _, d_row in grup.iterrows():
                        d_id = d_row["id"]
                        d_bolge = d_row.get("bolge", "-")
                        d_sistem = d_row.get("sistem_adi", "-")
                        d_parca = d_row.get("parca_adi", "-")
                        d_sn = d_row.get("parca_sn", "-")
                        d_pn = d_row.get("parca_pn", "-")
                        d_tarih = d_row.get("ekleme_tarihi", "-")
                        d_aciklama = d_row.get("aciklama", "")
                        
                        col_chk, col_card, col_acts = st.columns([0.6, 7.6, 1.8])
                        
                        with col_chk:
                            if st.checkbox("Seç", key=f"chk_depo_{d_id}", label_visibility="collapsed"):
                                secilen_idler.append(d_id)

                        with col_card:
                            aciklama_html = f"<div style='font-size:12px; color:#888; margin-top:4px;'>📝 Not: {d_aciklama}</div>" if pd.notna(d_aciklama) and str(d_aciklama).strip() else ""
                            bolge_badge = f"<span class='tag-badge' style='background-color:#1e242b; color:#00adb5; border:1px solid #00adb5;'>📍 {d_bolge}</span>" if d_bolge and d_bolge != "-" else ""
                            
                            st.markdown(
                                f"""
                                <div class="item-card" style='margin-bottom: 5px; padding: 8px 12px;'>
                                    <div class="item-title" style='font-size: 14px;'>📦 {d_parca} <span class="tag-badge status-sevk">SEVK EDİLDİ</span> {bolge_badge}</div>
                                    <div class="item-sub" style='font-size: 12px;'>
                                        <b>Sistem:</b> {d_sistem} | <b>SN:</b> <code style='color:#00adb5;'>{d_sn}</code> | <b>PN:</b> {d_pn} | <b>Tarih:</b> {d_tarih}
                                    </div>
                                    {aciklama_html}
                                </div>
                                """,
                                unsafe_allow_html=True,
                            )
                        
                        with col_acts:
                            col_sil = st.columns(1)[0]
                            with col_sil:
                                if st.button("🗑️", key=f"depo_sil_{d_id}", help="Depodan Sil", use_container_width=True):
                                    with get_db_connection() as conn:
                                        with conn.cursor() as cursor:
                                            cursor.execute("DELETE FROM depo_stok WHERE id=%s", (d_id,))
                                        conn.commit()
                                    st.cache_data.clear()
                                    log_yaz("DEPO SİLME", f"Depodan ID {d_id} ({d_parca}) silindi.")
                                    st.toast("✅ Kayıt silindi.", icon="🗑️")
                                    st.rerun()

            if secilen_idler:
                st.markdown("<br>", unsafe_allow_html=True)
                col_info, col_btn = st.columns([6, 4])
                col_info.warning(f"⚠️ Seçilen **{len(secilen_idler)}** adet malzemeyi kalıcı olarak silmek üzeresiniz.")
                if col_btn.button("🗑️ Seçilenleri Toplu Sil", type="primary", use_container_width=True):
                    with get_db_connection() as conn:
                        with conn.cursor() as cursor:
                            if len(secilen_idler) == 1:
                                cursor.execute("DELETE FROM depo_stok WHERE id = %s", (secilen_idler[0],))
                            else:
                                cursor.execute("DELETE FROM depo_stok WHERE id IN %s", (tuple(secilen_idler),))
                        conn.commit()
                    st.cache_data.clear()
                    log_yaz("TOPLU DEPO SİLME", f"{len(secilen_idler)} adet malzeme depodan toplu olarak silindi.")
                    st.toast(f"✅ Seçilen {len(secilen_idler)} malzeme silindi!", icon="🗑️")
                    st.rerun()
        else:
            st.success("Depoda şu an sevk edilmiş malzeme bulunmuyor. 🎉")

except Exception as e:
    st.error(f"Kritik Hata: {e}")
    st.code(traceback.format_exc())
